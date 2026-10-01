import io
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE_SOURCE_NOTES = {
    'jake': "Jake's Resume LaTeX source supplied by the project owner; rendered as an editable DOCX with its compact single-column hierarchy and section rules.",
    'faangpath': 'FAANGPath resume source supplied by the project owner; rendered as an editable DOCX with strong section separators and aligned metadata.',
    'awesome': 'Awesome-CV source supplied by the project owner; rendered as an editable DOCX interpretation preserving its strong visual hierarchy.',
    'rendercv': 'RenderCV source supplied by the project owner; rendered as an editable DOCX interpretation with structured sections.',
    'deedy': 'RoleFit two-column technical layout based on the supplied source-template family.',
    'classic': 'RoleFit ATS-first single-column layout with ruled section headers.',
    'modern': 'RoleFit modern ATS-first layout with aligned header metadata and section rules.',
    'engineering': 'RoleFit engineering layout with compact technical sidebar and aligned main content.',
    'alta': 'RoleFit professional layout with a strong header and ruled sections.'
}


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def _set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def _set_table_borders(table, color='777777', size='5', inside=False):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)

    names = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    for name in names:
        element = borders.find(qn('w:' + name))
        if element is None:
            element = OxmlElement('w:' + name)
            borders.append(element)
        if name.startswith('inside') and not inside:
            element.set(qn('w:val'), 'nil')
        else:
            element.set(qn('w:val'), 'single')
            element.set(qn('w:sz'), size)
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), color)


def _remove_table_borders(table):
    _set_table_borders(table, 'FFFFFF', '0', False)
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is not None:
        for child in borders:
            child.set(qn('w:val'), 'nil')


def _set_font(run, name, size, color=None, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _bottom_border(paragraph, color='777777', size='6', space='1'):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn('w:pBdr'))
    if pbdr is None:
        pbdr = OxmlElement('w:pBdr')
        pPr.append(pbdr)
    bottom = pbdr.find(qn('w:bottom'))
    if bottom is None:
        bottom = OxmlElement('w:bottom')
        pbdr.append(bottom)
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color)


def _add_bullet(doc, text, font, size=9.2, color='333333', indent=.18):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.first_line_indent = Inches(-.12)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run('• ' + text.lstrip('•- '))
    _set_font(r, font, size, color)
    return p


def _base_doc(font='Arial', size=9.5, margins=(.45, .45, .55, .55)):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(margins[0])
    sec.bottom_margin = Inches(margins[1])
    sec.left_margin = Inches(margins[2])
    sec.right_margin = Inches(margins[3])
    normal = doc.styles['Normal']
    normal.font.name = font
    normal.font.size = Pt(size)
    normal.font.color.rgb = RGBColor.from_string('333333')
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    return doc


def _header(doc, name, role, contact, font, accent, align=WD_ALIGN_PARAGRAPH.CENTER, name_size=21, ruled=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(name)
    _set_font(r, font, name_size, accent, True)

    if role:
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(role)
        _set_font(r, font, 10, '444444', True)

    if contact:
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(' | '.join(contact))
        _set_font(r, font, 8.4, '555555')

    if ruled:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _bottom_border(p, accent if accent != '111111' else '555555', '8', '0')


def _section_heading(doc, heading, font, accent='111111', size=9.3, rule=True, upper=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    text = heading.upper() if upper else heading
    r = p.add_run(text)
    _set_font(r, font, size, accent, True)
    if rule:
        _bottom_border(p, accent if accent != '111111' else '666666', '6', '1')
    return p


def _section(doc, heading, items, font, accent, size=9.2, compact=False, rule=True):
    if not items:
        return
    _section_heading(doc, heading, font, accent, size, rule)
    for item in items:
        if item.startswith(('•', '-')):
            _add_bullet(doc, item, font, size - .2)
        elif ' | ' in item:
            parts = [x.strip() for x in item.split(' | ') if x.strip()]
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1.5)
            r = p.add_run(parts[0])
            _set_font(r, font, size - .15, '222222', True)
            if len(parts) > 1:
                r = p.add_run(' | ' + ' | '.join(parts[1:]))
                _set_font(r, font, size - .15, '333333')
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1.5)
            r = p.add_run(item)
            _set_font(r, font, size - .15, '333333')


def _parse(text):
    lines = [x.strip() for x in (text or '').splitlines() if x.strip()]
    name = lines[0] if lines else 'Candidate'
    role = lines[1] if len(lines) > 1 else ''
    headings = {
        'SUMMARY', 'SKILLS', 'EXPERIENCE', 'EDUCATION', 'PROJECTS',
        'CERTIFICATIONS', 'ACHIEVEMENTS', 'OPTIMIZATION NOTES',
        'WORK EXPERIENCE', 'PROFESSIONAL EXPERIENCE', 'TECHNICAL SKILLS',
        'LEADERSHIP / EXTRACURRICULAR'
    }
    contact = []
    data = []
    current = None
    for line in lines[2:]:
        u = line.upper().rstrip(':')
        if u in headings:
            current = 'EXPERIENCE' if u in ('WORK EXPERIENCE', 'PROFESSIONAL EXPERIENCE') else u
            data.append((current, []))
            continue
        if current:
            data[-1][1].append(line)
        elif len(contact) < 2:
            contact.append(line)
    return name, role, contact, data


def _add_link_line(doc, text, font, size=8.4):
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    _set_font(r, font, size, '555555')


def _two_col_sections(doc, left_data, right_data, font, accent='111111', left_width=2.05, right_width=4.7, left_fill='F4F5F7'):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(left_width)
    table.columns[1].width = Inches(right_width)
    left, right = table.rows[0].cells
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_margins(cell, 70, 110, 70, 110)
    _set_table_borders(table, 'D2D6DC', '5', False)
    _shade(left, left_fill)

    def fill(cell, data, size):
        first = True
        for heading, items in data:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(heading.upper())
            _set_font(r, font, size, accent, True)
            _bottom_border(p, 'B7BDC7', '4', '1')
            for item in items:
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(('• ' + item.lstrip('•- ')) if item.startswith(('•', '-')) else item)
                _set_font(r, font, size - .5, '333333')

    fill(left, left_data, 8.6)
    fill(right, right_data, 9.0)
    return table


def _header_table(doc, name, role, contact, font, accent, left_width=5.4, right_width=1.25):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(left_width)
    table.columns[1].width = Inches(right_width)
    left, right = table.rows[0].cells
    _remove_table_borders(table)
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell, 20, 50, 20, 50)
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(name)
    _set_font(r, font, 22, accent, True)
    if role:
        p = left.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(role)
        _set_font(r, font, 10, '555555', True)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('\n'.join(contact))
    _set_font(r, font, 8.2, '555555')
    return table


def build_docx(text, template):
    name, role, contact, data = _parse(text)

    if template == 'jake':
        # Jake-style: centered identity, compact Times hierarchy and strong horizontal rules.
        doc = _base_doc('Times New Roman', 9.15, (.40, .40, .58, .58))
        _header(doc, name, role, contact, 'Times New Roman', '111111', WD_ALIGN_PARAGRAPH.CENTER, 20, ruled=True)
        for heading, items in data:
            _section(doc, heading, items, 'Times New Roman', '111111', 9.1, True, True)

    elif template == 'faangpath':
        doc = _base_doc('Arial', 9.1, (.40, .40, .55, .55))
        _header(doc, name, role, contact, 'Arial', '111111', WD_ALIGN_PARAGRAPH.CENTER, 19, ruled=True)
        for heading, items in data:
            _section(doc, heading, items, 'Arial', '111111', 9.0, True, True)

    elif template == 'awesome':
        doc = _base_doc('Calibri', 9.35, (.48, .48, .62, .62))
        _header(doc, name, role, contact, 'Calibri', '1F6F9B', WD_ALIGN_PARAGRAPH.LEFT, 22, ruled=True)
        for heading, items in data:
            _section(doc, heading, items, 'Calibri', '1F6F9B', 9.2, False, True)

    elif template == 'rendercv':
        doc = _base_doc('Aptos', 9.25, (.46, .46, .62, .62))
        _header(doc, name, role, contact, 'Aptos', '111827', WD_ALIGN_PARAGRAPH.LEFT, 21, ruled=True)
        for heading, items in data:
            _section(doc, heading, items, 'Aptos', '374151', 9.1, False, True)

    elif template == 'deedy':
        doc = _base_doc('Arial', 9.0, (.40, .40, .48, .48))
        _header(doc, name, role, contact, 'Arial', '111111', WD_ALIGN_PARAGRAPH.CENTER, 19, ruled=True)
        left_data = data[::2]
        right_data = data[1::2]
        _two_col_sections(doc, left_data, right_data, 'Arial', '111111', 2.15, 4.6, 'F4F5F7')

    elif template == 'engineering':
        doc = _base_doc('Consolas', 8.7, (.35, .35, .48, .48))
        _header(doc, name, role, contact, 'Consolas', '14532D', WD_ALIGN_PARAGRAPH.LEFT, 17, ruled=True)
        left = [x for x in data if x[0] in ('SKILLS', 'CERTIFICATIONS', 'EDUCATION')]
        right = [x for x in data if x[0] not in ('SKILLS', 'CERTIFICATIONS', 'EDUCATION')]
        _two_col_sections(doc, left, right, 'Consolas', '14532D', 2.15, 4.55, 'EEF6F0')

    elif template == 'classic':
        doc = _base_doc('Georgia', 9.45, (.48, .48, .62, .62))
        _header(doc, name, role, contact, 'Georgia', '222222', WD_ALIGN_PARAGRAPH.CENTER, 20, ruled=True)
        for heading, items in data:
            _section(doc, heading, items, 'Georgia', '222222', 9.25, False, True)

    elif template == 'modern':
        doc = _base_doc('Aptos', 9.65, (.52, .52, .65, .65))
        _header_table(doc, name, role, contact, 'Aptos', '4F46E5')
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _bottom_border(p, '4F46E5', '10', '0')
        for heading, items in data:
            _section(doc, heading, items, 'Aptos', '4F46E5', 9.45, False, True)

    else:  # alta / fallback
        doc = _base_doc('Calibri', 9.4, (.48, .48, .62, .62))
        _header(doc, name, role, contact, 'Calibri', '1976A8', WD_ALIGN_PARAGRAPH.LEFT, 21, ruled=True)
        for heading, items in data:
            _section(doc, heading, items, 'Calibri', '1976A8', 9.2, False, True)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
